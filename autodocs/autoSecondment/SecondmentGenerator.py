# ~*~ coding: utf-8 ~*~
import logging

import copy
import datetime
import json
import os

import docx
import openpyxl
from docx import Document
from openpyxl import load_workbook

from autoSecondment.generatorUtils.Constants import Constants as SecondmentConstants
from autoSecondment.generatorUtils.secondmentUtils import (createSecondmentDocument, indexParagrahps, indexRuns, mapCodeToCountry, calculateLastWorkingDate)
from autoSecondment.generatorUtils.PersonRecord import PersonRecord

class SecondmentGenerator:
    """
    Accepts secondment data, 
    and generate secondment docx files
    """
    currentMonth = None
    fileName = None
    personRecords = []
    
    def __init__(self):
        pass

    def setPersonRecords(self, personRecords):
        self.personRecords = personRecords

    def generateSecondments(self):
        # Constants for placeholders, template name, mappings and parameters.
        c = SecondmentConstants()

        # Open template .docx file, to use in each following file creation
        doc = Document(c.TEMPLATE_NAME)
        try:
            os.mkdir(c.DOCUMENTS_DIR)
        except FileExistsError:
            pass

        # Index all placeholder locations in the template .docx file
        paragraphIdx = indexParagrahps(c)
        runIdx = indexRuns(c, paragraphIdx)
        
        for record in self.personRecords:
            try:
                createSecondmentDocument(record.returnRecordArray(), copy.deepcopy(doc), paragraphIdx, runIdx, c, True)
            except Exception as e:
                logging.error("Could not create secondment; [SecondmentGenerator] error={}".format(e))
                logging.exception(e)
